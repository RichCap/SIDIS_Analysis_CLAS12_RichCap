#!/usr/bin/env python3
# Append the raw-A review slides to the existing deck. Existing slides stay in place.

import os

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

REVIEW = "/Users/richardcapobianco/Desktop/Work_Offline.nosync/SIDIS_Analysis_CLAS12_RichCap/Prepare_Next_Iteration/Fit_Refinement/plots/3D_BC_RC_acc020_review"
DECK = os.path.join(REVIEW, "3D_BC_RC_acc020_Final_Review.pptx")
MOSAIC = os.path.join(REVIEW, "mosaic")
COMPARE = os.path.join(REVIEW, "rawA_comparison")
SLIDE_PNG = os.path.join(MOSAIC, "for_slides")

INK = RGBColor(0x1C, 0x28, 0x33)
HEADER = RGBColor(0x1B, 0x3A, 0x4B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5E, 0x6A, 0x71)
CARD = RGBColor(0xFF, 0xFC, 0xF8)
BG = RGBColor(0xF7, 0xF5, 0xF2)


def render_pdf(pdf_path, png_path):
    document = fitz.open(pdf_path)
    pixmap = document[0].get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
    pixmap.save(png_path)
    document.close()
    return pixmap.width, pixmap.height


def set_run(paragraph, text, size, bold, color, font_name):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return run


def add_header(slide, title, slide_width):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), slide_width, Inches(0.72))
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADER
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.14), Inches(12.4), Inches(0.46))
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    set_run(paragraph, title, 24, True, WHITE, "Georgia")


def add_footer(slide, label, slide_width):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Inches(7.22), slide_width, Inches(0.28))
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADER
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.24), Inches(9.5), Inches(0.22))
    paragraph = box.text_frame.paragraphs[0]
    set_run(paragraph, label, 11, False, RGBColor(0xE7, 0xEE, 0xF2), "Calibri")
    right = slide.shapes.add_textbox(Inches(10.2), Inches(7.24), Inches(2.7), Inches(0.22))
    paragraph = right.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    set_run(paragraph, "3D BC+RC acc020", 11, False, RGBColor(0xE7, 0xEE, 0xF2), "Calibri")


def add_body(slide, x, y, w, h, lines):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for index, (text, size, bold, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if(index == 0) else frame.add_paragraph()
        paragraph.space_after = Pt(6)
        set_run(paragraph, text, size, bold, color, "Calibri")
    return box


def add_picture_fit(slide, path, x, y, max_w, max_h):
    from PIL import Image
    image = Image.open(path)
    aspect = image.size[0] / float(image.size[1])
    width = max_w
    height = width / aspect
    if(height > max_h):
        height = max_h
        width = height * aspect
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(width), Inches(height))
    return width, height


def mosaic_slide(prs, blank, title, png_path, caption):
    slide = prs.slides.add_slide(blank)
    add_header(slide, title, prs.slide_width)
    add_picture_fit(slide, png_path, 0.35, 0.88, 8.5, 6.15)
    add_body(slide, 9.05, 1.15, 3.85, 5.6, [
        (caption, 15, False, INK),
    ])
    add_footer(slide, "Measured points and the continuous spline.", prs.slide_width)
    return slide


def main():
    os.makedirs(SLIDE_PNG, exist_ok=True)
    pages = {
        "normA": os.path.join(MOSAIC, "acc020_normA_3D_BayesianUnfold_BC_RC_z_NormalizedA.pdf"),
        "rawA": os.path.join(MOSAIC, "acc020_rawA_3D_BayesianUnfold_BC_RC_z_Amplitude.pdf"),
        "B": os.path.join(MOSAIC, "acc020_B_3D_BayesianUnfold_BC_RC_z_CosPhiMoment.pdf"),
        "C": os.path.join(MOSAIC, "acc020_C_3D_BayesianUnfold_BC_RC_z_Cos2PhiMoment.pdf"),
    }
    pngs = {}
    for key, pdf_path in pages.items():
        png_path = os.path.join(SLIDE_PNG, key + ".png")
        render_pdf(pdf_path, png_path)
        pngs[key] = png_path
    prs = Presentation(DECK)
    before = len(prs.slides)
    def slide_text(one_slide):
        parts = []
        for shape in one_slide.shapes:
            if(shape.has_text_frame):
                parts.append(shape.text_frame.text)
        return "\n".join(parts)
    first_text = slide_text(prs.slides[0])
    blank = prs.slide_layouts[len(prs.slide_layouts) - 1]
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Two A splines, one pair of B and C models", prs.slide_width)
    add_body(slide, 0.5, 1.05, 12.3, 5.8, [
        ("Cross-section A", 18, True, HEADER),
        ("acc020_3D_BC_RC_..._Compute_SplineWeight.txt divides A by bin area, luminosity, and the virtual-photon flux. Use this file for cross sections. B and C are not rescaled.", 16, False, INK),
        ("Monte Carlo A", 18, True, HEADER),
        ("acc020_3D_BC_RC_rawA_..._Compute_SplineWeight.txt uses the fit amplitude directly. No luminosity, bin volume, beam-blocker, tracking, or photon-flux factor is applied. This is the new event-weight default.", 16, False, INK),
        ("Shared modulation", 18, True, HEADER),
        ("B and C in the raw-A file are copies of the existing 3D models, not a second fit. The July 5D weight file remains in the scripts as a commented line. Its absolute scale is a historical reference, not the target for these weights.", 16, False, INK),
    ])
    add_footer(slide, "Histogram fits were not redone.", prs.slide_width)
    mosaic_slide(prs, blank, "Normalized A versus z", pngs["normA"], "Logarithmic A0 after bin area, luminosity, and photon flux. Each pad is one Q2-y bin. Colors are pT bins. Dashed curves are the normalized-A spline.")
    mosaic_slide(prs, blank, "Unnormalized A versus z", pngs["rawA"], "The same fits before those factors. The vertical scale is the raw amplitude, about 10^4, so it is not drawn on the normalized axis.")
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Do the two A splines agree after normalization?", prs.slide_width)
    add_body(slide, 0.45, 1.0, 6.3, 5.9, [
        ("452 bin centers. Raw-spline A is divided by area x L x flux, then compared with the spline that was fit to already-normalized A.", 15, False, INK),
        ("Median signed difference: -0.35%.", 16, True, INK),
        ("Median absolute difference: 10.5%.", 16, True, INK),
        ("16th to 84th percentile: -15% to +17%.", 16, False, INK),
        ("95th percentile of the absolute difference: 42%.", 16, False, INK),
        ("Largest: Q2-y 8, z-pT 24, +79%. Nearby (8, 29) is +75%.", 16, False, INK),
        ("The raw spline matches the JSON amplitudes (median ratio 0.994). The extra spread is the log-space fit seeing a kinematic scale that changes from bin to bin. Both splines are kept.", 15, False, INK),
    ])
    add_picture_fit(slide, os.path.join(COMPARE, "relative_difference.png"), 6.9, 0.95, 6.0, 3.15)
    add_picture_fit(slide, os.path.join(COMPARE, "median_abs_rel_by_q2y.png"), 6.9, 4.2, 6.0, 2.8)
    add_footer(slide, "Common normalized convention. No bin was below the relative floor.", prs.slide_width)
    mosaic_slide(prs, blank, "B versus z", pngs["B"], "cos(phi_h) coefficient. The curve is the existing B spline, shared by both A families.")
    mosaic_slide(prs, blank, "C versus z", pngs["C"], "cos(2 phi_h) coefficient. The curve is the existing C spline, shared by both A families.")
    out_path = DECK + ".tmp"
    prs.save(out_path)
    check = Presentation(out_path)
    after_first = slide_text(check.slides[0])
    print("slides", before, "->", len(check.slides))
    print("first_same", first_text == after_first)
    print("first", after_first[:80].replace("\n", " "))
    if((len(check.slides) != before + 6) or (first_text != after_first) or (os.path.getsize(out_path) < 15000000)):
        raise SystemExit("deck append did not preserve the original slides")
    os.replace(out_path, DECK)
    print("replaced", DECK, os.path.getsize(DECK))


if(__name__ == "__main__"):
    main()
